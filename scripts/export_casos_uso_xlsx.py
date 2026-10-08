#!/usr/bin/env python3
"""
Script para generar o actualizar el libro Excel de Casos de Uso
basado en la plantilla oficial de la Universidad de Cartagena.
Uso:
    python export_casos_uso_xlsx.py --input casos_de_uso.json --output "Plantilla Casos de Uso Generada.xlsx"
"""

import os
import sys
import json
import argparse
from copy import copy
import openpyxl

def generate_use_cases_workbook(data_file: str, template_file: str, output_file: str):
    with open(data_file, 'r', encoding='utf-8') as f:
        data = json.load(f)

    if not os.path.exists(template_file):
        raise FileNotFoundError(f"No se encontró la plantilla base: {template_file}")

    wb = openpyxl.load_workbook(template_file)
    base_sheet = wb.active

    use_cases = data if isinstance(data, list) else data.get("casos_de_uso", [data])

    for i, cu in enumerate(use_cases):
        sheet_name = f"{cu.get('id', f'CU-{i+1}')}"
        if i == 0:
            ws = base_sheet
            ws.title = sheet_name
        else:
            ws = wb.copy_worksheet(base_sheet)
            ws.title = sheet_name

        # Mapeo de datos básicos
        ws['A1'] = f"Caso de Uso: {cu.get('nombre', 'Caso de uso')}"
        ws['A3'] = f"Id: {cu.get('id', f'CU-{i+1:02d}')}"
        ws['A4'] = "Breve descripcion:"
        ws['A5'] = cu.get('breve_descripcion', '')

        # Actores
        ws['A6'] = "Actor Principal:"
        principales = cu.get('actor_principal', [])
        ws['A7'] = "\n".join([f"{idx+1}. {a}" for idx, a in enumerate(principales)]) if isinstance(principales, list) else f"1. {principales}"

        ws['A9'] = "Actores Secundarios:"
        secundarios = cu.get('actores_secundarios', [])
        ws['A10'] = "\n".join([f"{idx+1}. {a}" for idx, a in enumerate(secundarios)]) if isinstance(secundarios, list) else f"1. {secundarios}"

        # Requisitos funcionales
        ws['A12'] = "Requisitos Funcionales Asociados:"
        rfs = cu.get('requisitos_funcionales_asociados', [])
        ws['A13'] = "\n".join([f"{idx+1}. {rf}" for idx, rf in enumerate(rfs)]) if isinstance(rfs, list) else f"1. {rfs}"

        # Objetivos
        ws['A15'] = "Objetivos:"
        objs = cu.get('objetivos', [])
        ws['A16'] = "\n".join([f"{idx+1}. {o}" for idx, o in enumerate(objs)]) if isinstance(objs, list) else f"1. {objs}"

        # Precondiciones
        ws['A18'] = "Precondiciones:"
        preconds = cu.get('precondiciones', [])
        if isinstance(preconds, list):
            ws['A19'] = "\n".join([f"{idx+1}. {p}" for idx, p in enumerate(preconds)])
        else:
            ws['A19'] = f"1. {preconds}"

        # Datos requeridos
        ws['A22'] = "Datos requeridos:"
        datos = cu.get('datos_requeridos', [])
        ws['A23'] = "\n".join([f"{idx+1}. {d}" for idx, d in enumerate(datos)]) if isinstance(datos, list) else f"1. {datos}"

        # Flujo principal
        ws['A25'] = "Flujo principal:"
        flujo = cu.get('flujo_principal', [])
        flujo_lines = []
        for idx, paso in enumerate(flujo):
            if isinstance(paso, dict):
                flujo_lines.append(f"{paso.get('paso', idx+1)}. ({paso.get('actor', 'Sistema')}) {paso.get('accion', '')}")
            else:
                flujo_lines.append(f"{idx+1}. {paso}")
        ws['A26'] = "\n".join(flujo_lines)

        # Postcondiciones
        ws['A32'] = "Postcondiciones:"
        posts = cu.get('postcondiciones', [])
        ws['A33'] = "\n".join([f"{idx+1}. {p}" for idx, p in enumerate(posts)]) if isinstance(posts, list) else f"1. {posts}"

    wb.save(output_file)
    print(f"Libro de casos de uso generado exitosamente en: {output_file}")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Generador de Excel para Casos de Uso UdeC")
    parser.add_argument("--input", required=True, help="Archivo JSON con la definición de casos de uso")
    parser.add_argument("--template", default=os.path.join(os.path.dirname(__file__), "..", "templates", "plantilla_casos_uso.xlsx"), help="Ruta a plantilla base")
    parser.add_argument("--output", default="Casos_de_Uso.xlsx", help="Ruta del archivo Excel generado")
    args = parser.parse_args()

    generate_use_cases_workbook(args.input, args.template, args.output)
